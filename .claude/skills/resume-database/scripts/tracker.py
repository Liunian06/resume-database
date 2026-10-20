#!/usr/bin/env python3
"""
投递记录表管理工具

用于管理候选人的投递记录表（application_tracker.xlsx）。
"""

import argparse
import sys
from pathlib import Path
from typing import Optional

try:
    import openpyxl
    from openpyxl import Workbook
    from openpyxl.styles import Font
except ImportError:
    print("ERROR: 需要安装openpyxl库：pip install openpyxl", file=sys.stderr)
    sys.exit(1)


COLUMNS = ["序号", "企业", "岗位名称", "岗位JD", "进度", "投递链接", "备注"]
STATUS_OPTIONS = ["待投递", "已投递", "简历筛选中", "已约面试", "面试中", "已拒绝", "已通过", "已放弃"]


def get_tracker_path(user_id: str) -> Path:
    """获取投递记录表路径"""
    return Path(f"data/{user_id}/application_tracker.xlsx")


def init_tracker(user_id: str) -> None:
    """初始化投递记录表"""
    tracker_path = get_tracker_path(user_id)

    if tracker_path.exists():
        print(f"ERROR: 投递记录表已存在：{tracker_path}")
        sys.exit(1)

    user_dir = Path(f"data/{user_id}")
    if not user_dir.exists():
        print(f"ERROR: 用户目录不存在：{user_dir}")
        sys.exit(1)

    wb = Workbook()
    ws = wb.active
    ws.title = "投递记录"

    # 写入表头
    for col_idx, col_name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = Font(bold=True)

    wb.save(tracker_path)
    print(f"投递记录表已创建：{tracker_path}")


def add_record(
    user_id: str,
    company: str,
    position: str,
    jd_path: str,
    status: str = "待投递",
    link: str = "",
    note: str = ""
) -> None:
    """添加新投递记录"""
    tracker_path = get_tracker_path(user_id)

    if not tracker_path.exists():
        print("ERROR: 投递记录表不存在，请先运行 tracker.py init")
        sys.exit(1)

    wb = openpyxl.load_workbook(tracker_path)
    ws = wb.active

    # 计算新序号
    next_id = ws.max_row  # 表头占第1行，数据从第2行开始

    # 添加新行
    ws.append([next_id, company, position, jd_path, status, link, note])

    wb.save(tracker_path)

    print(f"已添加投递记录：")
    print(f"序号: {next_id}")
    print(f"企业: {company}")
    print(f"岗位名称: {position}")
    print(f"进度: {status}")


def update_record(
    user_id: str,
    record_id: int,
    status: Optional[str] = None,
    link: Optional[str] = None,
    note: Optional[str] = None
) -> None:
    """更新投递记录"""
    tracker_path = get_tracker_path(user_id)

    if not tracker_path.exists():
        print("ERROR: 投递记录表不存在")
        sys.exit(1)

    wb = openpyxl.load_workbook(tracker_path)
    ws = wb.active

    # 查找记录（record_id对应行号 = record_id + 1，因为表头占第1行）
    row_idx = record_id + 1

    if row_idx > ws.max_row:
        print(f"ERROR: 投递记录 #{record_id} 不存在")
        sys.exit(1)

    if status is None and link is None and note is None:
        print("ERROR: 请至少提供一个更新字段（status/link/note）")
        sys.exit(1)

    # 更新字段
    if status is not None:
        ws.cell(row=row_idx, column=5, value=status)
    if link is not None:
        ws.cell(row=row_idx, column=6, value=link)
    if note is not None:
        ws.cell(row=row_idx, column=7, value=note)

    wb.save(tracker_path)

    # 读取更新后的记录
    company = ws.cell(row=row_idx, column=2).value
    position = ws.cell(row=row_idx, column=3).value
    current_status = ws.cell(row=row_idx, column=5).value
    current_note = ws.cell(row=row_idx, column=7).value

    print(f"已更新投递记录 #{record_id}：")
    print(f"企业: {company}")
    print(f"岗位名称: {position}")
    print(f"进度: {current_status}")
    if current_note:
        print(f"备注: {current_note}")


def query_records(
    user_id: str,
    company: Optional[str] = None,
    status: Optional[str] = None,
    position: Optional[str] = None
) -> None:
    """查询投递记录"""
    tracker_path = get_tracker_path(user_id)

    if not tracker_path.exists():
        print("ERROR: 投递记录表不存在")
        sys.exit(1)

    wb = openpyxl.load_workbook(tracker_path)
    ws = wb.active

    results = []

    # 遍历所有记录（跳过表头）
    for row_idx in range(2, ws.max_row + 1):
        record_id = ws.cell(row=row_idx, column=1).value
        record_company = ws.cell(row=row_idx, column=2).value
        record_position = ws.cell(row=row_idx, column=3).value
        record_jd = ws.cell(row=row_idx, column=4).value
        record_status = ws.cell(row=row_idx, column=5).value
        record_link = ws.cell(row=row_idx, column=6).value or ""
        record_note = ws.cell(row=row_idx, column=7).value or ""

        # 应用过滤条件
        if company and company not in str(record_company):
            continue
        if status and status != record_status:
            continue
        if position and position not in str(record_position):
            continue

        results.append({
            "id": record_id,
            "company": record_company,
            "position": record_position,
            "jd": record_jd,
            "status": record_status,
            "link": record_link,
            "note": record_note
        })

    print(f"查询结果（{len(results)}条）：")
    if not results:
        return

    print()
    for record in results:
        print(f"#{record['id']} | {record['company']} | {record['position']} | {record['status']}")
        print(f"JD: {record['jd']}")
        if record['link']:
            print(f"链接: {record['link']}")
        if record['note']:
            print(f"备注: {record['note']}")
        print()


def list_records(user_id: str) -> None:
    """列出所有投递记录"""
    tracker_path = get_tracker_path(user_id)

    if not tracker_path.exists():
        print("ERROR: 投递记录表不存在")
        sys.exit(1)

    wb = openpyxl.load_workbook(tracker_path)
    ws = wb.active

    total = ws.max_row - 1  # 减去表头
    print(f"投递记录列表（{total}条）：")

    if total == 0:
        return

    print()
    for row_idx in range(2, ws.max_row + 1):
        record_id = ws.cell(row=row_idx, column=1).value
        company = ws.cell(row=row_idx, column=2).value
        position = ws.cell(row=row_idx, column=3).value
        status = ws.cell(row=row_idx, column=5).value
        print(f"#{record_id} | {company} | {position} | {status}")


def main():
    parser = argparse.ArgumentParser(
        description="投递记录表管理工具",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
示例:
  python tracker.py init --user-id 10001
  python tracker.py add --user-id 10001 --company "字节跳动" --position "产品实习生" --jd "data/10001/jobs/bytedance-product/jd.md"
  python tracker.py update --user-id 10001 --id 1 --status "已投递"
  python tracker.py query --user-id 10001 --status "已投递"
  python tracker.py list --user-id 10001
"""
    )

    subparsers = parser.add_subparsers(dest="command", help="子命令")

    # init子命令
    init_parser = subparsers.add_parser("init", help="初始化投递记录表")
    init_parser.add_argument("--user-id", required=True, help="用户ID")

    # add子命令
    add_parser = subparsers.add_parser("add", help="添加新投递记录")
    add_parser.add_argument("--user-id", required=True, help="用户ID")
    add_parser.add_argument("--company", required=True, help="公司名称")
    add_parser.add_argument("--position", required=True, help="岗位名称")
    add_parser.add_argument("--jd", required=True, help="JD文件路径")
    add_parser.add_argument("--status", default="待投递", help="进度状态")
    add_parser.add_argument("--link", default="", help="投递链接")
    add_parser.add_argument("--note", default="", help="备注")

    # update子命令
    update_parser = subparsers.add_parser("update", help="更新投递记录")
    update_parser.add_argument("--user-id", required=True, help="用户ID")
    update_parser.add_argument("--id", type=int, required=True, help="记录序号")
    update_parser.add_argument("--status", help="进度状态")
    update_parser.add_argument("--link", help="投递链接")
    update_parser.add_argument("--note", help="备注")

    # query子命令
    query_parser = subparsers.add_parser("query", help="查询投递记录")
    query_parser.add_argument("--user-id", required=True, help="用户ID")
    query_parser.add_argument("--company", help="公司名称（模糊匹配）")
    query_parser.add_argument("--status", help="进度状态（精确匹配）")
    query_parser.add_argument("--position", help="岗位名称（模糊匹配）")

    # list子命令
    list_parser = subparsers.add_parser("list", help="列出所有投递记录")
    list_parser.add_argument("--user-id", required=True, help="用户ID")

    args = parser.parse_args()

    if not args.command:
        parser.print_help()
        sys.exit(0)

    if args.command == "init":
        init_tracker(args.user_id)
    elif args.command == "add":
        add_record(
            args.user_id,
            args.company,
            args.position,
            args.jd,
            args.status,
            args.link,
            args.note
        )
    elif args.command == "update":
        update_record(
            args.user_id,
            args.id,
            args.status,
            args.link,
            args.note
        )
    elif args.command == "query":
        query_records(
            args.user_id,
            args.company,
            args.status,
            args.position
        )
    elif args.command == "list":
        list_records(args.user_id)


if __name__ == "__main__":
    main()