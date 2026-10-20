#!/usr/bin/env python3
"""
把简历 Markdown 转成 DOCX。Markdown 格式见 assets/templates/resume_output_template.md，
与本脚本一一对应，模板之外的写法直接报错。
"""

from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BASIC_INFO_TITLE = "基本信息"
BASIC_INFO_FIELDS = {
    "姓名": "name",
    "求职意向": "target_role",
    "手机": "phone",
    "邮箱": "email",
    "求职城市": "city",
    "证件照": "photo",
}
EDUCATION_HEADER = ["学校", "学历", "专业", "时间"]
BODY_COLUMN_WIDTHS_CM = (7.3, 2.7, 5.0, 4.0)
LEADING_LABEL_RE = re.compile(r"^\*\*([^*\r\n]+)\*\*\s*(.*)$")

PHOTO_WIDTH_CM = 2.2
PHOTO_HEIGHT_CM = 3.1
PHOTO_RIGHT_INSET_CM = 0.5  # 照片距页面右侧正文边线的距离
HEADER_SIDE_COLUMN_CM = PHOTO_WIDTH_CM + PHOTO_RIGHT_INSET_CM + 0.1  # 0.1 余量防止图片被单元格边距挤出
HEADER_COLUMN_WIDTHS_CM = (HEADER_SIDE_COLUMN_CM, 19.0 - 2 * HEADER_SIDE_COLUMN_CM, HEADER_SIDE_COLUMN_CM)
DOCX_FONT_NAME = "MiSans"
NAME_FONT_SIZE = 18
META_FONT_SIZE = 10
BODY_FONT_SIZE = 10
SECTION_TITLE_FONT_SIZE = 11
ENTRY_TITLE_FONT_SIZE = 10


class ResumeFormatError(ValueError):
    pass


@dataclass
class ExperienceEntry:
    name: str
    role: str
    date: str
    bullets: List[str] = field(default_factory=list)


@dataclass
class EducationBlock:
    rows: List[List[str]] = field(default_factory=list)
    bullets: List[str] = field(default_factory=list)


@dataclass
class Section:
    title: str
    kind: str  # "table" | "experience" | "list"
    education: List[EducationBlock] = field(default_factory=list)
    entries: List[ExperienceEntry] = field(default_factory=list)
    bullets: List[str] = field(default_factory=list)


@dataclass
class Resume:
    label: str
    meta: Dict[str, str]
    sections: List[Section]
    base_dir: Path



# ---------------------------------------------------------------- 解析 ----

def fail(line_no: int, message: str) -> "ResumeFormatError":
    return ResumeFormatError(f"第{line_no}行：{message}")


def split_sections(lines: Sequence[str]) -> Tuple[str, List[Tuple[str, List[Tuple[int, str]]]]]:
    label: Optional[str] = None
    sections: List[Tuple[str, List[Tuple[int, str]]]] = []
    for line_no, raw in enumerate(lines, start=1):
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# "):
            if label is not None:
                raise fail(line_no, "一级标题只能有一个")
            if sections:
                raise fail(line_no, "一级标题必须在所有板块之前")
            label = line[2:].strip()
            continue
        if line.startswith("## "):
            sections.append((line[3:].strip(), []))
            continue
        if not sections:
            raise fail(line_no, "一级标题之后必须先出现 `## 基本信息`")
        sections[-1][1].append((line_no, line))
    if not label:
        raise ResumeFormatError("缺少一级标题 `# <企业>-<岗位>-<姓名>`")
    return label, sections


def parse_basic_info(body: Sequence[Tuple[int, str]]) -> Dict[str, str]:
    meta: Dict[str, str] = {}
    for line_no, line in body:
        if not line.startswith("- ") or "：" not in line:
            raise fail(line_no, "基本信息只能写 `- 键：值`")
        key, value = (part.strip() for part in line[2:].split("：", 1))
        if key not in BASIC_INFO_FIELDS:
            raise fail(line_no, f"基本信息不支持字段“{key}”，只支持：{'、'.join(BASIC_INFO_FIELDS)}")
        if key == "手机" and value and not re.fullmatch(r"\d{11}", value):
            raise fail(line_no, "手机必须是11位纯数字，不加-或空格")
        if value:
            meta[BASIC_INFO_FIELDS[key]] = value
    if not meta.get("name"):
        raise ResumeFormatError("基本信息缺少“姓名”")
    return meta


def parse_table_row(line_no: int, line: str) -> List[str]:
    if not (line.startswith("|") and line.endswith("|")):
        raise fail(line_no, "表格行必须以 | 开头并以 | 结尾")
    return [cell.strip() for cell in line[1:-1].split("|")]


def parse_table_section(title: str, body: Sequence[Tuple[int, str]]) -> Section:
    """教育经历有两种互斥写法，按结构自动识别，同一份简历不能混用：
    分段式：每段学历一个表格块（表头+分隔行+1行数据），其后的 `- ` 行属于这一段学历，块可重复多次
    合并式：只有一个表格块，里面有多行学历，其后一组 `- ` 行是总描述
    """
    section = Section(title=title, kind="table")
    block: Optional[EducationBlock] = None
    has_bullets = False  # 当前块是否已进入描述部分
    for line_no, line in body:
        if line.startswith("|"):
            cells = parse_table_row(line_no, line)
            if cells == EDUCATION_HEADER:  # 表头永远开启一个新块
                block = EducationBlock()
                section.education.append(block)
                has_bullets = False
            elif block is None:
                raise fail(line_no, f"表格必须先写表头 | {' | '.join(EDUCATION_HEADER)} |")
            elif all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                continue
            elif has_bullets:
                raise fail(line_no, "`- ` 行之后要继续写学历，必须先写新的表头")
            elif len(cells) != len(EDUCATION_HEADER):
                raise fail(line_no, f"表格每行必须是{len(EDUCATION_HEADER)}列")
            else:
                block.rows.append(cells)
        elif line.startswith("- "):
            if block is None or not block.rows:
                raise fail(line_no, "`- ` 行之前必须先有学历表格")
            block.bullets.append(line[2:].strip())
            has_bullets = True
        else:
            raise fail(line_no, "表格板块只能有表格行和 `- ` 行")

    if not section.education or any(not b.rows for b in section.education):
        raise ResumeFormatError(f"“{title}”板块有表格没有数据行")
    if len(section.education) > 1 and any(len(b.rows) > 1 for b in section.education):
        raise ResumeFormatError(f"“{title}”不能混用两种写法：要么每段学历一个表格，要么所有学历放进同一个表格")
    return section


def parse_experience_section(title: str, body: Sequence[Tuple[int, str]]) -> Section:
    section = Section(title=title, kind="experience")
    current: Optional[ExperienceEntry] = None
    for line_no, line in body:
        if line.startswith("### "):
            parts = [part.strip() for part in line[4:].split("｜")]
            if len(parts) != 3 or not all(parts):
                raise fail(line_no, "经历标题必须是 `### 名称｜职位｜时间`（全角｜，三段都不能空）")
            current = ExperienceEntry(*parts)
            section.entries.append(current)
        elif line.startswith("- "):
            if current is None:
                raise fail(line_no, "`- ` 行之前必须先有 `### 名称｜职位｜时间`")
            current.bullets.append(line[2:].strip())
        else:
            raise fail(line_no, "经历板块只能有 `###` 标题和 `- ` 行")
    return section


def parse_list_section(title: str, body: Sequence[Tuple[int, str]]) -> Section:
    section = Section(title=title, kind="list")
    for line_no, line in body:
        if not line.startswith("- "):
            raise fail(line_no, "列表板块只能有 `- ` 行")
        section.bullets.append(line[2:].strip())
    return section


def parse_resume_markdown(path: Path) -> Resume:
    label, raw_sections = split_sections(path.read_text(encoding="utf-8").splitlines())
    if not raw_sections or raw_sections[0][0] != BASIC_INFO_TITLE:
        raise ResumeFormatError(f"第一个板块必须是 `## {BASIC_INFO_TITLE}`")
    meta = parse_basic_info(raw_sections[0][1])

    sections: List[Section] = []
    for title, body in raw_sections[1:]:
        if title == BASIC_INFO_TITLE:
            raise ResumeFormatError("基本信息只能出现一次")
        if any(line.startswith("|") for _, line in body):
            sections.append(parse_table_section(title, body))
        elif any(line.startswith("### ") for _, line in body):
            sections.append(parse_experience_section(title, body))
        else:
            sections.append(parse_list_section(title, body))
    if meta.get("photo") and not (path.parent / meta["photo"]).is_file():
        raise ResumeFormatError(f"证件照文件不存在：{path.parent / meta['photo']}")
    return Resume(label=label, meta=meta, sections=sections, base_dir=path.parent)



# ---------------------------------------------------------------- 版式 ----

def set_all_fonts(r_fonts) -> None:
    """中文、英文、数字、复杂文种全部指定为 MiSans，并去掉主题字体覆盖。"""
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attr}"), DOCX_FONT_NAME)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        key = qn(f"w:{attr}")
        if key in r_fonts.attrib:
            del r_fonts.attrib[key]


def set_run_font(run, size: float, bold: bool = False) -> None:
    run.font.name = DOCX_FONT_NAME
    set_all_fonts(run._element.rPr.rFonts)
    run.font.size = Pt(size)
    run.bold = bold


def zero_spacing(paragraph_format) -> None:
    paragraph_format.space_before = Pt(0)
    paragraph_format.space_after = Pt(0)
    paragraph_format.line_spacing = 1.0


def define_style(document: Document, name: str, size: float, bold: bool = False,
                 align=None, base: Optional[str] = None, color: Optional[RGBColor] = None):
    style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    if base:
        style.base_style = document.styles[base]
    style.font.name = DOCX_FONT_NAME
    set_all_fonts(style._element.get_or_add_rPr().get_or_add_rFonts())
    style.font.size = Pt(size)
    style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color
    if align is not None:
        style.paragraph_format.alignment = align
    zero_spacing(style.paragraph_format)
    return style


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.0)
    section.left_margin = section.right_margin = Cm(1.0)
    doc_grid = section._sectPr.find(qn("w:docGrid"))
    if doc_grid is not None:
        section._sectPr.remove(doc_grid)

    # 默认模板的内置样式带主题字体，会盖过 MiSans
    for r_fonts in document.styles.element.iter(qn("w:rFonts")):
        set_all_fonts(r_fonts)

    normal = document.styles["Normal"]
    normal.font.name = DOCX_FONT_NAME
    set_all_fonts(normal._element.get_or_add_rPr().get_or_add_rFonts())
    normal.font.size = Pt(BODY_FONT_SIZE)
    zero_spacing(normal.paragraph_format)

    define_style(document, "Resume Name", NAME_FONT_SIZE, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    define_style(document, "Resume Meta", META_FONT_SIZE, align=WD_ALIGN_PARAGRAPH.CENTER,
                 color=RGBColor(80, 80, 80))
    define_style(document, "Resume Section", SECTION_TITLE_FONT_SIZE, bold=True).paragraph_format.keep_with_next = True
    bullet = define_style(document, "Resume Bullet", BODY_FONT_SIZE, base="List Bullet")
    bullet.paragraph_format.left_indent = Cm(0.55)
    bullet.paragraph_format.first_line_indent = Cm(-0.28)
    define_style(document, "Resume Table Text", BODY_FONT_SIZE)


def set_cell_bottom_border(cell, size_eighth_pt: int = 10) -> None:
    """sz 单位为 1/8 磅，10 = 1.25 磅。"""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    bottom = OxmlElement("w:bottom")
    for key, value in (("val", "single"), ("sz", str(size_eighth_pt)), ("space", "0"), ("color", "000000")):
        bottom.set(qn(f"w:{key}"), value)
    borders.append(bottom)
    anchor = None
    for tag in ("w:vMerge", "w:gridSpan", "w:tcW"):
        anchor = tc_pr.find(qn(tag))
        if anchor is not None:
            break
    if anchor is not None:
        anchor.addnext(borders)
    else:
        tc_pr.insert(0, borders)


def set_table_layout(table, widths_cm: Sequence[float]) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        borders.append(element)
    cell_mar = OxmlElement("w:tblCellMar")
    for side, value in (("top", 0), ("left", 20), ("bottom", 0), ("right", 20)):
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")
        cell_mar.append(element)
    # tblPr 子元素顺序：tblW, tblBorders, tblLayout, tblCellMar, tblLook
    tbl_pr.find(qn("w:tblLayout")).addprevious(borders)
    tbl_pr.find(qn("w:tblLayout")).addnext(cell_mar)
    for column, width in zip(table.columns, widths_cm):
        column.width = Cm(width)


def add_row(table, widths_cm: Sequence[float] = BODY_COLUMN_WIDTHS_CM):
    row = table.add_row()
    for cell, width in zip(row.cells, widths_cm):
        cell.width = Cm(width)
    return row


def fill_cell(cell, text: str, align, size: float, bold: bool) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.style = "Resume Table Text"
    paragraph.alignment = align
    set_run_font(paragraph.add_run(text), size, bold)


def add_bullet_paragraph(cell, text: str, first: bool) -> None:
    paragraph = cell.paragraphs[0] if first else cell.add_paragraph()
    paragraph.style = "Resume Bullet"
    paragraph.paragraph_format.keep_together = True
    match = LEADING_LABEL_RE.match(text)
    if match:
        set_run_font(paragraph.add_run(match.group(1)), BODY_FONT_SIZE, bold=True)
        set_run_font(paragraph.add_run(match.group(2)), BODY_FONT_SIZE)
    else:
        set_run_font(paragraph.add_run(text), BODY_FONT_SIZE)


# ---------------------------------------------------------------- 输出 ----

LEFT, RIGHT = WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT


def add_header_lines(container, meta: Dict[str, str], first_paragraph=None) -> None:
    """container 是 Document 或单元格；有证件照时写进中间单元格，复用同一套文字。"""
    def paragraph(style: str):
        nonlocal first_paragraph
        if first_paragraph is not None:
            p, first_paragraph = first_paragraph, None
            p.style = style
            return p
        return container.add_paragraph(style=style)

    set_run_font(paragraph("Resume Name").add_run(meta["name"]), NAME_FONT_SIZE, bold=True)
    for label, key in (("求职意向", "target_role"), ("求职城市", "city")):
        if meta.get(key):
            set_run_font(paragraph("Resume Meta").add_run(f"{label}：{meta[key]}"), META_FONT_SIZE)
    contact = []
    if meta.get("phone"):
        phone = meta["phone"]
        contact.append(f"{phone[:3]} {phone[3:7]} {phone[7:]}")
    if meta.get("email"):
        contact.append(meta["email"])
    if contact:
        set_run_font(paragraph("Resume Meta").add_run(" | ".join(contact)), META_FONT_SIZE)


def add_header(document: Document, resume: Resume) -> None:
    meta = resume.meta
    if not meta.get("photo"):
        add_header_lines(document, meta)
        return
    photo_path = resume.base_dir / meta["photo"]
    # 三列：左侧留白与照片等宽，姓名等文字仍居页面正中，右侧放照片
    table = document.add_table(rows=1, cols=3)
    set_table_layout(table, HEADER_COLUMN_WIDTHS_CM)
    left, middle, right = table.rows[0].cells
    for cell, width in zip((left, middle, right), HEADER_COLUMN_WIDTHS_CM):
        cell.width = Cm(width)
    middle.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    add_header_lines(middle, meta, first_paragraph=middle.paragraphs[0])
    right_mar = OxmlElement("w:tcMar")
    inset = OxmlElement("w:right")
    inset.set(qn("w:w"), str(int(round(PHOTO_RIGHT_INSET_CM / 2.54 * 1440))))
    inset.set(qn("w:type"), "dxa")
    right_mar.append(inset)
    right._tc.get_or_add_tcPr().append(right_mar)
    photo_paragraph = right.paragraphs[0]
    photo_paragraph.style = "Resume Table Text"
    photo_paragraph.alignment = RIGHT
    photo_paragraph.add_run().add_picture(str(photo_path), width=Cm(PHOTO_WIDTH_CM), height=Cm(PHOTO_HEIGHT_CM))


def merged_row(table):
    row = add_row(table)
    return row.cells[0].merge(row.cells[3])


def add_section_title_row(table, title: str) -> None:
    cell = merged_row(table)
    paragraph = cell.paragraphs[0]
    paragraph.style = "Resume Section"
    set_run_font(paragraph.add_run(title), SECTION_TITLE_FONT_SIZE, bold=True)
    set_cell_bottom_border(cell)


def add_bullet_block_row(table, bullets: Sequence[str]) -> None:
    if not bullets:
        return
    cell = merged_row(table)
    for index, bullet in enumerate(bullets):
        add_bullet_paragraph(cell, bullet, first=index == 0)


def add_education_rows(table, section: Section) -> None:
    for block in section.education:
        for values in block.rows:
            row = add_row(table)
            for cell, value, align in zip(row.cells, values, (LEFT, LEFT, LEFT, RIGHT)):
                fill_cell(cell, value, align, BODY_FONT_SIZE, bold=True)
        add_bullet_block_row(table, block.bullets)


def add_experience_rows(table, section: Section) -> None:
    for entry in section.entries:
        row = add_row(table)
        name_cell = row.cells[0].merge(row.cells[1])  # 前两列合并，形成 10.0 / 5.0 / 4.0 三列
        for cell, value, align in zip((name_cell, row.cells[2], row.cells[3]),
                                      (entry.name, entry.role, entry.date), (LEFT, LEFT, RIGHT)):
            fill_cell(cell, value, align, ENTRY_TITLE_FONT_SIZE, bold=True)
        add_bullet_block_row(table, entry.bullets)


def build_docx(resume: Resume, output_docx: Path) -> None:
    document = Document()
    configure_document(document)
    add_header(document, resume)
    table = document.add_table(rows=0, cols=len(BODY_COLUMN_WIDTHS_CM))
    set_table_layout(table, BODY_COLUMN_WIDTHS_CM)
    for section in resume.sections:
        add_section_title_row(table, section.title)
        if section.kind == "table":
            add_education_rows(table, section)
        elif section.kind == "experience":
            add_experience_rows(table, section)
        else:
            add_bullet_block_row(table, section.bullets)
    output_docx.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_docx)


# ------------------------------------------------------------ PDF（可选）----

def find_soffice() -> Optional[str]:
    for command in ("soffice.com", "soffice", "soffice.exe"):
        resolved = shutil.which(command)
        if resolved:
            return resolved
    for candidate in (r"C:\Program Files\LibreOffice\program\soffice.com",
                      r"C:\Program Files (x86)\LibreOffice\program\soffice.com"):
        if Path(candidate).exists():
            return candidate
    return None


def export_pdf(docx_path: Path, timeout: int) -> Path:
    soffice = find_soffice()
    if not soffice:
        raise RuntimeError("未找到 LibreOffice，无法导出 PDF")
    result = subprocess.run(
        [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(docx_path.parent), str(docx_path)],
        capture_output=True, text=True, check=False, timeout=timeout,
    )
    pdf_path = docx_path.with_suffix(".pdf")
    if result.returncode != 0 or not pdf_path.exists():
        raise RuntimeError(f"LibreOffice 导出失败：{(result.stderr or result.stdout).strip()}")
    return pdf_path


# ---------------------------------------------------------------- 入口 ----

def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = argparse.ArgumentParser(description="把简历 Markdown 转成 DOCX。")
    parser.add_argument("--input", type=Path, required=True, help="简历 Markdown 路径")
    parser.add_argument("--out-dir", type=Path, default=None, help="输出目录，默认与输入文件同目录")
    parser.add_argument("--docx", type=str, default=None, help="DOCX 文件名，默认与输入文件同名")
    parser.add_argument("--pdf", action="store_true", help="额外导出 PDF（需要 LibreOffice），默认不导出")
    parser.add_argument("--pdf-timeout", type=int, default=120, help="PDF 导出超时秒数")
    args = parser.parse_args(argv)

    input_path = args.input.resolve()
    out_dir = args.out_dir.resolve() if args.out_dir else input_path.parent
    docx_path = out_dir / (args.docx or f"{input_path.stem}.docx")

    try:
        resume = parse_resume_markdown(input_path)
    except ResumeFormatError as exc:
        print(f"ERROR: Markdown 格式错误：{exc}", file=sys.stderr)
        return 1
    build_docx(resume, docx_path)
    print(f"DOCX created: {docx_path}")

    if args.pdf:
        try:
            print(f"PDF created: {export_pdf(docx_path, args.pdf_timeout)}")
        except Exception as exc:
            print(f"PDF 导出失败（DOCX 已生成）：{exc}", file=sys.stderr)
            return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
